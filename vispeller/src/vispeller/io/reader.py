import json
import re
import time
import urllib.request
from io import BytesIO

from openpyxl import load_workbook
from tqdm import tqdm

GOOGLE_SHEETS_RE = re.compile(r"docs\.google\.com/spreadsheets/d/([a-zA-Z0-9-_]+)")

def normalize_link(link: str) -> str:
    match = GOOGLE_SHEETS_RE.search(link)
    if not match:
        return link
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=xlsx"

class SpreadsheetReader:
    def __init__(self, link: str, name: str = None):
        self.link = normalize_link(link)
        self.name = name or link
        self._workbook = None

    def _get_workbook(self):
        if self._workbook is None:
            print(f"downloading {self.name}...")
            t0 = time.time()
            with urllib.request.urlopen(self.link, timeout=15) as response:
                total = int(response.headers.get("Content-Length", 0))
                chunks = []
                with tqdm(total=total, unit="B", unit_scale=True, desc="Downloading") as bar:
                    while True:
                        chunk = response.read(65536)
                        if not chunk:
                            break
                        chunks.append(chunk)
                        bar.update(len(chunk))
                raw_bytes = b"".join(chunks)
            print(f"downloaded {len(raw_bytes)} bytes in {time.time() - t0:.2f}s")

            print("parsing workbook...")
            t1 = time.time()
            self._workbook = load_workbook(
                filename=BytesIO(raw_bytes),
                data_only=True,
                read_only=True,
            )
            print(f"parsed in {time.time() - t1:.2f}s")
            if self._workbook.properties.title:
                self.name = self._workbook.properties.title
        return self._workbook

    def sheet_names(self):
        return self._get_workbook().sheetnames

    def fetch(self, sheet_index: int = 0, show_progress: bool = True) -> str:
        workbook = self._get_workbook()
        sheet_names = workbook.sheetnames

        if sheet_index < 0 or sheet_index >= len(sheet_names):
            raise IndexError(
                f"sheet_index={sheet_index} is out of range, "
                f"workbook only has {len(sheet_names)} sheet (0 -> {len(sheet_names) - 1})"
            )

        sheet_name = sheet_names[sheet_index]
        sheet = workbook[sheet_name]

        cells = []
        row_iter = sheet.iter_rows(values_only=True)

        if show_progress:
            row_iter = tqdm(
                row_iter,
                total=sheet.max_row,
                desc=f"Reading '{sheet_name}'",
                unit="row",
                leave=False,
            )

        for row_idx, row in enumerate(row_iter, start=1):
            for col_idx, value in enumerate(row, start=1):
                if value is not None and str(value).strip() != "":
                    cells.append({
                        "row": row_idx,
                        "column": col_idx,
                        "text": str(value),
                    })

        return json.dumps({sheet_name: cells}, ensure_ascii=False)